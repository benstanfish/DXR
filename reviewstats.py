# Copyright (c) 2018-2025 Ben Fisher

import sys, os

from openpyxl import Workbook
from openpyxl.styles import DEFAULT_FONT
from openpyxl.worksheet.formula import ArrayFormula


from PyQt6.QtWidgets import QApplication, QFileDialog

from dxbuild.reviews import Review
from dxbuild.constants import FALLBACKS, _LOG_DIR
from dxbuild.buildtools import timestamp, clean_name, autoincrement_name
from dxreport import singlereport

if not os.path.exists(_LOG_DIR):
    os.makedirs(_LOG_DIR)

import logging
from dxcore.logconstants import log_format_string
logger = logging.getLogger(__name__)
logger.setLevel(logging.WARNING)
log_formatter = logging.Formatter(log_format_string)
log_file_handler = logging.FileHandler(f'{_LOG_DIR}/{__name__}.log')
log_file_handler.setFormatter(log_formatter)
logger.addHandler(log_file_handler)



def main() -> None:

    _WRITE_FILE = True
    xml_path = './dev/test/data.xml'

    DEFAULT_FONT.__init__(name=FALLBACKS['font_name'], size=FALLBACKS['font_size'])
    wb = Workbook()

    # app = QApplication(sys.argv)
    # xml_paths, _ = QFileDialog.getOpenFileNames(
    #     parent=None, 
    #     caption='Select Files Dialog',
    #     filter='XML Files (*.xml);;HTML Files (*.html);;All Files (*)'
    # )

    review = Review.from_file(xml_path)
    ws = wb.active
    
    singlereport.create_report(review, ws)

    ws_table_name = ws.tables.items()[0][0]

    ws_stat_name = ws.title[:25] + '-STAT'
    ws_stats = wb.create_sheet(ws_stat_name)
    
    comments = review.review_comments.comments
    disciplines = list(set([comment.discipline for comment in comments]))
    for disc in disciplines:
        print(disc, len(disciplines))

    ws_stats['A1'].value = f'=UNIQUE({ws_table_name}[Discipline])'




    if _WRITE_FILE:
        save_name = f'./dev/test/out/test_{timestamp()}.xlsx'
        wb.save(save_name)
        logger.debug(f'_WRITE_FILE = {_WRITE_FILE} -> saved workbook to "{save_name}"')
    wb.close()


if __name__ == '__main__':
    main()
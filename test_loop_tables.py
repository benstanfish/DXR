
import openpyxl

wb = openpyxl.load_workbook('./dev/test/tables_test.xlsx')

def get_table_names(wb: Workbook) -> list:
    table_names = []
    try:

        for ws in wb.worksheets:
            if ws.tables:
                for table_name, table_obj in ws.tables.items():
                    table_names.append(table_name)
    finally:
        wb.close()
    return table_names

